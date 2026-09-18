"""Le calendrier 2026 de Colorplast tel que le cabinet le tient, salarié par salarié.

Source : `data/colorplast/variables/2026-08/calendrier-2026-colorplast.xlsx`
(déposé par Gaëlle sur le Drive, un onglet par salarié, douze blocs de quatre
colonnes : jour, quantième, « H.Abs », « CP »). C'est le calendrier
d'absences de l'entreprise, pas le bulletin : congés payés (CP = 1, 0,5),
fériés non payés (« JFNP »), journées de récupération (« EN RECUP »), absence
(« absente »), événement familial (« EVF »), jour d'accident du travail, et,
pour un mois d'entrée, les heures faites jour par jour (Demory, mars).
Les fériés y sont nommés dans « H.Abs » pour tout le monde ; ce ne sont pas
des absences.

Usage : python -m scripts.backtest.colorplast_calendrier_cabinet [MOIS_MAX]
"""

from __future__ import annotations

import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl

CLASSEUR = Path(__file__).resolve().parents[3] / "data" / "colorplast" / "variables" / "2026-08" / "calendrier-2026-colorplast.xlsx"
MOIS = ["JANVIER", "FEVRIER", "MARS", "AVRIL", "MAI", "JUIN", "JUILLET", "AOUT", "SEPTEMBRE", "OCTOBRE", "NOVEMBRE", "DECEMBRE"]
FERIES = ("jour de l'an", "lundi de pâques", "fete du w", "victoire 1945", "ascension", "pentecote", "fête nationale",
          "assomption", "toussaint", "armistice 1918", "noël")
SALARIES = ("BUGNY", "COTTE", "ESPINOSA", "FUCKAR", "GAUTHERON", "GIRERD", "DEMORY")


@dataclass(frozen=True)
class Evenement:
    salarie: str
    mois: int
    jour: int
    nature: str        # cp | jfnp | recup | absence | evenement_familial | accident_travail | heures | formation | visite_medicale | note
    valeur: float | str | None
    brut: str          # ce qui est écrit dans le classeur


def _nature(habs, cp) -> list[tuple[str, float | str | None, str]]:
    out = []
    h = "" if habs in (None, " ", "") else str(habs).strip()
    c = None if cp in (None, " ", "") else cp
    hl = h.lower()
    if isinstance(habs, (int, float)) and habs not in (0,):
        out.append(("heures", float(habs), h))
    elif h and hl not in FERIES:
        if "recup" in hl:
            out.append(("recup", 1.0, h))
        elif "absente" in hl or "pas travaill" in hl:
            out.append(("absence", 1.0, h))
        elif hl.startswith("evf"):
            out.append(("evenement_familial", 1.0, h))
        elif "at" in hl.split() or "jour at" in hl:
            out.append(("accident_travail", 1.0, h))
        elif "formation" in hl:
            out.append(("formation", None, h))
        elif hl.startswith("vm"):
            out.append(("visite_medicale", None, h))
        elif hl == "cp":
            pass  # la colonne CP porte la quantité
        else:
            out.append(("note", None, h))
    if c is not None:
        if isinstance(c, (int, float)):
            out.append(("cp", float(c), str(c)))
        else:
            cl = str(c).strip().lower()
            if cl == "jfnp":
                out.append(("jfnp", 1.0, str(c)))
            elif "recup" in cl:
                out.append(("recup", 1.0, str(c)))
            else:
                out.append(("note", None, str(c)))
    return out


def lire(mois_max: int = 12) -> list[Evenement]:
    wb = openpyxl.load_workbook(CLASSEUR, data_only=True)
    evenements: list[Evenement] = []
    for nom in SALARIES:
        ws = wb[nom]
        colonnes = {str(c.value).strip().upper(): c.column for c in ws[1] if c.value and str(c.value).strip().upper() in MOIS}
        for libelle, col in colonnes.items():
            mois = MOIS.index(libelle) + 1
            if mois > mois_max:
                continue
            for r in range(2, ws.max_row + 1):
                quantieme = ws.cell(r, col + 1).value
                if not isinstance(quantieme, (int, float)):
                    continue
                for nature, valeur, brut in _nature(ws.cell(r, col + 2).value, ws.cell(r, col + 3).value):
                    evenements.append(Evenement(nom, mois, int(quantieme), nature, valeur, brut))
    return evenements


if __name__ == "__main__":
    mois_max = int(sys.argv[1]) if len(sys.argv) > 1 else 12
    for ev in lire(mois_max):
        print(f"{ev.salarie:10s} {ev.jour:02d}/{ev.mois:02d}  {ev.nature:18s} {ev.valeur!s:6s}  « {ev.brut} »")
